from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .services import AssetRecord, build_report_bytes, calculate_schedule


class DepreciationServiceTests(SimpleTestCase):
    def assertMoney(self, actual, expected):
        actual = Decimal(str(actual)).quantize(Decimal("0.01"))
        expected = Decimal(str(expected)).quantize(Decimal("0.01"))
        self.assertEqual(actual, expected)

    def test_straight_line_matches_current_year_sample_policy(self):
        asset = AssetRecord(
            no=1,
            description="Sample asset",
            account_code="1601106",
            unit="1 pcs",
            purchase_date=date(2025, 2, 5),
            useful_life=48,
            asset_amount=Decimal("375000"),
            depreciation_method="Straight Line",
        )
        beginning, months, total, ending = calculate_schedule(asset, 2025)
        self.assertMoney(beginning, "375000")
        self.assertMoney(months[0], "0")
        self.assertMoney(months[1], "7812.50")
        self.assertMoney(total, "85937.50")
        self.assertMoney(ending, "289062.50")

    def test_ddb_prorated_example(self):
        asset = AssetRecord(
            no=1,
            description="DDB prorated example",
            account_code="1601",
            unit="1 unit",
            purchase_date=date(2024, 7, 30),
            useful_life=48,
            asset_amount=Decimal("4000000"),
            depreciation_method="DDB",
        )
        beginning_2024, months_2024, total_2024, ending_2024 = calculate_schedule(asset, 2024)
        self.assertMoney(beginning_2024, "4000000")
        self.assertMoney(total_2024, "1000000")
        self.assertMoney(ending_2024, "3000000")
        self.assertMoney(sum(months_2024[6:12]), "1000000")

        beginning_2025, months_2025, total_2025, ending_2025 = calculate_schedule(asset, 2025)
        self.assertMoney(beginning_2025, "3000000")
        self.assertMoney(total_2025, "1500000")
        self.assertMoney(ending_2025, "1500000")
        self.assertMoney(sum(months_2025), "1500000")

    def test_build_report_from_sample_upload(self):
        sample = Path(__file__).resolve().parent.parent / "samples" / "Sampel upload file.xlsx"
        if not sample.exists():
            self.skipTest("Sample upload file is not available.")
        output, metadata = build_report_bytes(sample, client_name="Transit Pratama Utama", year=2025)
        self.assertEqual(metadata.row_count, 4)
        wb = load_workbook(BytesIO(output.getvalue()), data_only=True)
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Client ")
        self.assertEqual(ws["C1"].value, ":  Transit Pratama Utama")
        self.assertEqual(ws["I6"].value, 375000)
        self.assertEqual(ws["J6"].value, 0)
        self.assertMoney(ws["K6"].value, "7812.50")
        self.assertMoney(ws["V6"].value, "85937.50")
        self.assertMoney(ws["W6"].value, "289062.50")
