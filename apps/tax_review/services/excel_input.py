from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import re

from openpyxl import load_workbook

from .constants import FINANCIAL_REQUIRED_HEADERS, MAPPING_REQUIRED_HEADERS, OBJECT_KEYS
from .exceptions import InputWorkbookError


@dataclass(frozen=True)
class FinancialRow:
    account_name: object
    description: str
    amount: float
    source_row: int


@dataclass(frozen=True)
class KeywordRule:
    keyword: str
    sign: int
    flags: dict
    source_row: int


def _header_map(ws):
    values = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    return {value: idx for idx, value in enumerate(values)}


def _require_headers(ws, required, file_label):
    header = _header_map(ws)
    missing = [name for name in required if name not in header]
    if missing:
        raise InputWorkbookError(f"{file_label}: kolom wajib tidak ditemukan: {', '.join(missing)}")
    return header


def _to_number(value, context):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        raise InputWorkbookError(f"Nilai bukan angka pada {context}: {value!r}")


def read_keyword_rules(file_obj):
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    wb = load_workbook(file_obj, data_only=True, read_only=True, keep_links=False)
    ws = wb.worksheets[0]
    header = _require_headers(ws, MAPPING_REQUIRED_HEADERS, "Mapping Keyword")
    rules = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        keyword = row[header["Keyword"]]
        if keyword in (None, ""):
            continue
        sign_raw = row[header["Tanda"]]
        try:
            sign = int(sign_raw) if sign_raw not in (None, "") else 1
        except (TypeError, ValueError):
            raise InputWorkbookError(f"Mapping Keyword baris {row_no}: Tanda harus -1, 1, atau kosong.")
        if sign not in (-1, 1):
            raise InputWorkbookError(f"Mapping Keyword baris {row_no}: Tanda harus -1, 1, atau kosong.")
        flags = {}
        for key, col_name in OBJECT_KEYS.items():
            raw = row[header[col_name]]
            flags[key] = str(raw or "").strip().lower() in {"y", "yes", "1", "true", "x"}
        if not any(flags.values()):
            continue
        rules.append(KeywordRule(str(keyword).strip(), sign, flags, row_no))
    if not rules:
        raise InputWorkbookError("Mapping Keyword tidak memiliki rule aktif.")
    return rules


def infer_year(sheet_name):
    match = re.search(r"(20\d{2})", str(sheet_name))
    return int(match.group(1)) if match else None


def read_financial_workbook(file_obj):
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    wb = load_workbook(file_obj, data_only=True, read_only=True, keep_links=False)
    result = {}
    for ws in wb.worksheets:
        year = infer_year(ws.title)
        if not year:
            continue
        header = _require_headers(ws, FINANCIAL_REQUIRED_HEADERS, f"Laporan Keuangan sheet {ws.title}")
        rows = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in row):
                continue
            rows.append(FinancialRow(
                account_name=row[header["Nama Akun"]],
                description=str(row[header["Keterangan"]] or ""),
                amount=_to_number(row[header["Jumlah"]], f"sheet {ws.title} baris {row_no}"),
                source_row=row_no,
            ))
        result[year] = rows
    if not result:
        raise InputWorkbookError("Laporan Keuangan tidak memiliki sheet dengan nama tahun (contoh: 2024, 2025).")
    return result
