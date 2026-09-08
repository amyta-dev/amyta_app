import io
from copy import copy
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import OBJECT_KEYS

DARK = "1F4E78"
MID = "5B9BD5"
LIGHT = "D9EAF7"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="A6A6A6")
MONEY_FMT = '#,##0;[Red](#,##0);-'
MILLION_FMT = '#,##0.00;[Red](#,##0.00);-'
PERCENT_FMT = '0.00%'
DATE_FMT = 'dd-mmm-yyyy'


def _title(ws, company, subtitle, period_text=None, end_col=8):
    ws["A1"] = f"PT {company}" if not str(company).lower().startswith("pt ") else company
    ws["A2"] = subtitle
    if period_text:
        ws["A3"] = period_text
    for cell in ("A1", "A2"):
        ws[cell].font = Font(bold=True, size=12)
    ws.sheet_view.showGridLines = False


def _header_cell(cell, fill=DARK):
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=THIN)


def _style_data_area(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_rekap_simple(wb, sheet_name, company, period_title, parsed, header_row=5):
    ws = wb.create_sheet(sheet_name[:31])
    _title(ws, company, "Rekapitulasi SPT PPh Pasal 21" if "PPh 21" in sheet_name else "Rekapitulasi SPT PPh Unifikasi", period_title)
    headers = parsed["headers"]
    for col, header in enumerate(headers, start=1):
        c = ws.cell(header_row, col, header); _header_cell(c)
    for r_idx, row in enumerate(parsed["rows"], start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, value)
    _style_data_area(ws, header_row + 1, max(header_row + 1, ws.max_row), 1, len(headers))
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, ws.max_row)}"
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(max(12, len(str(headers[col-1])) * 0.65), 34)
    return ws


def add_rekap_ppn(wb, sheet_name, company, period_title, parsed, is_new):
    ws = wb.create_sheet(sheet_name[:31])
    _title(ws, company, "Rekapitulasi SPT PPN", period_title)
    headers = parsed["headers"]
    if is_new:
        groups = [(1,1,"Bulan"),(2,2,"Tahun"),(3,41,"Data SPT PPN"),(42,43,"Data BPE"),(44,46,"Data Pembayaran"),(47,47,"Catatan")]
    else:
        groups = [(1,1,"Bulan"),(2,2,"Tahun"),(3,16,"Data SPT PPN"),(17,18,"Data BPE"),(19,21,"Data Pembayaran"),(22,22,"Catatan")]
    for start,end,label in groups:
        if start != end:
            ws.merge_cells(start_row=5,start_column=start,end_row=5,end_column=end)
        ws.cell(5,start,label); _header_cell(ws.cell(5,start))
        for c in range(start+1,end+1):
            ws.cell(5,c).fill = PatternFill("solid",fgColor=DARK)
    for col, header in enumerate(headers, start=1):
        c=ws.cell(6,col,header if col not in (1,2) else None); _header_cell(c, MID)
    ws.merge_cells(start_row=5,start_column=1,end_row=6,end_column=1)
    ws.merge_cells(start_row=5,start_column=2,end_row=6,end_column=2)
    if headers[-1] == "Catatan":
        ws.merge_cells(start_row=5,start_column=len(headers),end_row=6,end_column=len(headers))
    for r_idx,row in enumerate(parsed["rows"],start=7):
        for c_idx,value in enumerate(row,start=1): ws.cell(r_idx,c_idx,value)
    _style_data_area(ws,7,max(7,ws.max_row),1,len(headers))
    ws.freeze_panes="C7"
    for col in range(1,len(headers)+1):
        width=14 if col<=3 else 22
        if col==len(headers): width=42
        ws.column_dimensions[get_column_letter(col)].width=width
    return ws


def add_mapped_financial_sheet(wb, company, abbreviation, period_label, year, rows):
    title=f"{abbreviation} {period_label}"[:31]
    ws=wb.create_sheet(title)
    period_desc = f"Periode Januari - Desember {year}" if period_label.startswith("FY") else f"Periode Januari - PTD {year}"
    _title(ws, company, "Kertas Kerja Review Perpajakan (Laporan Laba Rugi)", period_desc)
    headers=["Nama Akun","Keterangan","Jumlah","Objek PPh 23","Objek PPh 4(2)","Objek PPh 21","Objek PPN Keluaran",f"Koreksi {abbreviation}","Tambahan Koreksi Fiskal"]
    for i,h in enumerate(headers,1): _header_cell(ws.cell(5,i,h))
    for idx,row in enumerate(rows,start=6):
        ws.cell(idx,1,row.account_name); ws.cell(idx,2,row.description); ws.cell(idx,3,row.amount)
        for offset,key in enumerate(("pph23","pph4_2","pph21","ppn","koreksi_pph_badan"),start=4):
            value=row.values[key]
            if value is not None:
                # Formula keeps derivation transparent and follows BRD sign rule.
                sign = -1 if abs(value + row.amount) < 1e-7 else 1
                ws.cell(idx,offset, f"=-C{idx}" if sign == -1 else f"=C{idx}")
        ws.cell(idx,3).number_format=MONEY_FMT
        for c in range(4,10): ws.cell(idx,c).number_format=MONEY_FMT
    total_row=6+len(rows)+1
    ws.cell(total_row,2,"Total Objek Pajak").font=Font(bold=True)
    for c in range(4,10):
        ws.cell(total_row,c,f"=SUM({get_column_letter(c)}6:{get_column_letter(c)}{total_row-2})")
        ws.cell(total_row,c).font=Font(bold=True); ws.cell(total_row,c).number_format=MONEY_FMT
    ws.freeze_panes="A6"
    ws.auto_filter.ref=f"A5:I{max(5,total_row-2)}"
    widths=[16,42,18,18,18,18,20,18,22]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    return ws, total_row


def _sheet_ref(sheet_name):
    return "'" + str(sheet_name).replace("'", "''") + "'"


def _financial_object_formula(mapped_refs, year, column_letter):
    ref = mapped_refs.get(year)
    if not ref:
        return None
    return f"={_sheet_ref(ref['sheet'])}!{column_letter}{ref['total_row']}/1000000"


def _rekap_sum_formula(sheet, sum_col, year_col, year, extra_col=None, extra_criteria=None):
    sheet_ref = _sheet_ref(sheet)
    if extra_col:
        return (
            f'=SUMIFS({sheet_ref}!${sum_col}:${sum_col},{sheet_ref}!${year_col}:${year_col},{year},'
            f'{sheet_ref}!${extra_col}:${extra_col},"{extra_criteria}")/1000000'
        )
    return f'=SUMIFS({sheet_ref}!${sum_col}:${sum_col},{sheet_ref}!${year_col}:${year_col},{year})/1000000'


def _add_recon_formula_sheet(wb, name, company, period, description_rows, mapped_refs, source_formulas, rate, interest_total, admin_rate=None, note=""):
    ws = wb.create_sheet(name)
    subtitle = {
        "Rekon WHT 23 - Sewa Jasa": "Rekonsiliasi PPh Pasal 23 atas sewa dan jasa",
        "Rekon WHT 4(2) - Sewa": "Rekonsiliasi PPh Pasal 4(2) atas sewa tanah dan bangunan",
        "Rekon WHT 21": "Rekonsiliasi PPh Pasal 21",
        "Rekon PPN": "Rekonsiliasi PPN",
    }[name]
    _title(ws, company, subtitle)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)
    ws.cell(4, 1, "Deskripsi"); _header_cell(ws.cell(4, 1))
    start_col = 3
    for idx, label in enumerate(period.labels, start=start_col):
        _header_cell(ws.cell(4, idx, label))
    total_col = start_col + len(period.labels)
    _header_cell(ws.cell(4, total_col, "Total"))

    for r_idx, (label, code) in enumerate(description_rows, start=5):
        ws.cell(r_idx, 1, label); ws.cell(r_idx, 2, code)
        if r_idx in (5, 5 + len(description_rows) - 1):
            ws.cell(r_idx, 1).font = Font(bold=True)

    for year_idx, year in enumerate(period.years, start=start_col):
        if year not in mapped_refs:
            continue
        col = get_column_letter(year_idx)
        if name == "Rekon WHT 23 - Sewa Jasa":
            ws.cell(5, year_idx, _financial_object_formula(mapped_refs, year, "D"))
            ws.cell(6, year_idx, f"={col}5*{rate/100}")
            ws.cell(7, year_idx, source_formulas(year))
            ws.cell(8, year_idx, f"=MAX({col}6-{col}7,0)")
            ws.cell(9, year_idx, f"={col}8*{interest_total/100}")
            ws.cell(10, year_idx, f"={col}8+{col}9")
        elif name == "Rekon WHT 4(2) - Sewa":
            ws.cell(5, year_idx, _financial_object_formula(mapped_refs, year, "E"))
            ws.cell(6, year_idx, f"={col}5*{rate/100}")
            ws.cell(7, year_idx, source_formulas(year))
            ws.cell(8, year_idx, f"=MAX({col}6-{col}7,0)")
            ws.cell(9, year_idx, f"={col}8*{interest_total/100}")
            ws.cell(10, year_idx, f"={col}8+{col}9")
        elif name == "Rekon WHT 21":
            ws.cell(5, year_idx, _financial_object_formula(mapped_refs, year, "F"))
            ws.cell(6, year_idx, f"={col}5*{rate/100}")
            ws.cell(7, year_idx, source_formulas(year))
            ws.cell(8, year_idx, f"=MAX({col}6-{col}7,0)")
            ws.cell(9, year_idx, f"={col}8*{interest_total/100}")
            ws.cell(10, year_idx, f"={col}8+{col}9")
        else:
            ws.cell(5, year_idx, _financial_object_formula(mapped_refs, year, "G"))
            ws.cell(6, year_idx, source_formulas(year))
            ws.cell(7, year_idx, f"={col}5-{col}6")
            ws.cell(8, year_idx, f"=MAX({col}7*{rate/100},0)")
            ws.cell(9, year_idx, f"={col}8*{interest_total/100}")
            ws.cell(10, year_idx, f"=MAX({col}7*{(admin_rate or 0)/100},0)")
            ws.cell(11, year_idx, f"={col}8+{col}9+{col}10")

    last_data_row = 11 if name == "Rekon PPN" else 10
    first = get_column_letter(start_col); last = get_column_letter(total_col - 1)
    for r in range(5, last_data_row + 1):
        ws.cell(r, total_col, f"=SUM({first}{r}:{last}{r})")
        for c in range(start_col, total_col + 1):
            ws.cell(r, c).number_format = MILLION_FMT

    note_row = last_data_row + 3
    ws.cell(note_row, 1, note)
    ws.column_dimensions["A"].width = 54; ws.column_dimensions["B"].width = 20
    for c in range(start_col, total_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.sheet_view.showGridLines = False
    return ws


def add_reconciliation_sheets(wb, company, period, mapped_refs, sheet_names, params):
    interest_total = params.interest_total
    note = f"Sanksi bunga maksimal = {params.interest_monthly:g}% x {params.interest_months} bulan = {interest_total:g}%"

    uni = sheet_names["unifikasi"]
    def wht23_source(year):
        return _rekap_sum_formula(uni, "F", "B", year, "D", "*23*")
    def wht42_source(year):
        # Accept the wording used in the sample working paper.
        return _rekap_sum_formula(uni, "F", "B", year, "D", "*4 ayat 2*")

    p21_old, p21_new = sheet_names["pph21_old"], sheet_names["pph21_new"]
    def p21_source(year):
        if year <= 2024:
            return _rekap_sum_formula(p21_old, "H", "B", year)
        sh = _sheet_ref(p21_new)
        return f'=SUMIF({sh}!$A:$A,"*{year}*",{sh}!$C:$C)/1000000'

    ppn_old, ppn_new = sheet_names["ppn_old"], sheet_names["ppn_new"]
    def ppn_source(year):
        if year <= 2024:
            return _rekap_sum_formula(ppn_old, "I", "B", year)
        return _rekap_sum_formula(ppn_new, "L", "B", year)

    _add_recon_formula_sheet(
        wb, "Rekon WHT 23 - Sewa Jasa", company, period,
        [("Objek pemotongan menurut laporan keuangan", "a"),
         ("PPh Pasal 23 terutang atas sewa dan jasa", f"b = {params.wht23:g}% x a"),
         ("PPh Pasal 23 yang sudah disetor", "c"),
         ("Selisih kurang potong PPh Pasal 23", "d = b - c"),
         ("Sanksi bunga", f"e = {interest_total:g}% x d"),
         ("Total eksposur PPh 23 atas sewa dan jasa", "f = d + e")],
        mapped_refs, wht23_source, params.wht23, interest_total, note=note,
    )
    _add_recon_formula_sheet(
        wb, "Rekon WHT 4(2) - Sewa", company, period,
        [("Objek pemotongan menurut laporan keuangan", "a"),
         ("PPh Pasal 4(2) terutang atas sewa tanah dan bangunan", f"b = {params.wht4_2:g}% x a"),
         ("PPh Pasal 4(2) yang sudah disetor", "c"),
         ("Selisih kurang potong PPh Pasal 4(2)", "d = b - c"),
         ("Sanksi bunga", f"e = {interest_total:g}% x d"),
         ("Total eksposur PPh 4(2) atas sewa tanah dan bangunan", "f = d + e")],
        mapped_refs, wht42_source, params.wht4_2, interest_total, note=note,
    )
    _add_recon_formula_sheet(
        wb, "Rekon WHT 21", company, period,
        [("Objek pemotongan menurut laporan keuangan", "a"),
         ("PPh Pasal 21 terutang", f"b = {params.wht21:g}% x a"),
         ("PPh Pasal 21 yang sudah dilaporkan di SPT", "c"),
         ("Kurang bayar PPh 21", "d = b - c"),
         ("Sanksi bunga", f"e = {interest_total:g}% x d"),
         ("Total eksposur PPh 21", "f = d + e")],
        mapped_refs, p21_source, params.wht21, interest_total,
        note=f"Memakai tarif rata-rata PPh Pasal 21. {note}",
    )
    _add_recon_formula_sheet(
        wb, "Rekon PPN", company, period,
        [("Objek PPN keluaran menurut laporan keuangan", "a"),
         ("Objek PPN keluaran menurut SPT PPN", "b"),
         ("Selisih objek PPN", "c = a - b"),
         ("PPN kurang bayar", f"d = {params.ppn:g}% x c"),
         ("Sanksi bunga", f"e = {interest_total:g}% x d"),
         ("Sanksi administrasi", f"f = {params.ppn_admin:g}% x c"),
         ("Total eksposur PPN", "g = d + e + f")],
        mapped_refs, ppn_source, params.ppn, interest_total, admin_rate=params.ppn_admin, note=note,
    )

def build_workbook(company, abbreviation, period, parsed_spt, mapped_by_year, recon, params):
    wb = Workbook(); wb.remove(wb.active)
    old_sheet_suffix = f"FY{str(period.start_year)[-2:]}"
    later_sheet_suffix = (
        f"FY{str(period.start_year + 1)[-2:]}-PTD{str(period.end_year)[-2:]}"
        if period.end_year > period.start_year else f"PTD{str(period.end_year)[-2:]}"
    )
    old_period_title = f"FY{period.start_year}"
    later_period_title = (
        f"FY{period.start_year + 1} - PTD{period.end_year}"
        if period.end_year > period.start_year else f"PTD{period.end_year}"
    )
    sheet_names = {
        "pph21_old": f"Rekap PPh 21 {old_sheet_suffix}",
        "pph21_new": f"Rekap PPh 21 {later_sheet_suffix}",
        "ppn_old": f"Rekap PPN {old_sheet_suffix}",
        "ppn_new": f"Rekap PPN {later_sheet_suffix}",
        "unifikasi": f"Rekap PPh Unifikasi FY{str(period.start_year)[-2:]}-PTD{str(period.end_year)[-2:]}",
    }
    add_rekap_simple(wb, sheet_names["pph21_old"], company, old_period_title, parsed_spt["pph21_old"])
    add_rekap_simple(wb, sheet_names["pph21_new"], company, later_period_title, parsed_spt["pph21_new"])
    add_rekap_ppn(wb, sheet_names["ppn_old"], company, old_period_title, parsed_spt["ppn_old"], False)
    add_rekap_ppn(wb, sheet_names["ppn_new"], company, later_period_title, parsed_spt["ppn_new"], True)
    add_rekap_simple(
        wb, sheet_names["unifikasi"], company,
        f"FY{period.start_year} - PTD{period.end_year}", parsed_spt["unifikasi"],
    )

    mapped_refs = {}
    for year in period.years:
        rows = mapped_by_year.get(year, [])
        if not rows:
            continue
        ws, total_row = add_mapped_financial_sheet(
            wb, company, abbreviation, period.label_for_year(year), year, rows,
        )
        mapped_refs[year] = {"sheet": ws.title, "total_row": total_row}

    add_reconciliation_sheets(wb, company, period, mapped_refs, sheet_names, params)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    out = io.BytesIO(); wb.save(out); out.seek(0); return out
