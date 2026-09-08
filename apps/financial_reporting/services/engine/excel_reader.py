from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .constants import HEADER_ALIASES, TABLE_SPECS
from .models import InputTable, Issue
from .utils import clean_text, normalize_code, normalize_header, parse_date_value


class InputFormatError(ValueError):
    pass


def _header_lookup(spec: dict[str, Any]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for canonical in spec["columns"]:
        lookup[normalize_header(canonical)] = canonical
    for alias, canonical in HEADER_ALIASES.items():
        if canonical in spec["columns"]:
            lookup[normalize_header(alias)] = canonical
    return lookup


def _is_xlsx(path: Path) -> bool:
    try:
        with path.open("rb") as handle:
            return handle.read(4) == b"PK\x03\x04"
    except OSError:
        return False


def read_input_table(
    path: str | Path,
    *,
    spec_key: str,
    label: str,
    issues: list[Issue],
) -> InputTable:
    """Read one input workbook using a sequential streaming scan.

    Random calls to ``worksheet.cell`` are disproportionately expensive for
    ``read_only`` workbooks. A single forward pass keeps large or heavily
    formatted templates responsive and uses bounded memory on Render.
    """

    path = Path(path)
    if path.suffix.casefold() != ".xlsx" or not _is_xlsx(path):
        raise InputFormatError(f"{label}: file harus berformat .xlsx yang valid.")

    spec = TABLE_SPECS[spec_key]
    lookup = _header_lookup(spec)
    required = set(spec["required"])

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise InputFormatError(f"{label}: workbook tidak dapat dibuka ({exc}).") from exc

    try:
        best_observed: set[str] = set()
        for worksheet in workbook.worksheets:
            max_scan_col = min(max(worksheet.max_column or 1, 1), 100)
            row_iterator = worksheet.iter_rows(min_row=1, max_col=max_scan_col)
            header_row = 0
            reverse: dict[str, int] = {}
            rows: list[dict[str, Any]] = []
            keep_blank_key_rows = bool(spec.get("keep_blank_key_rows"))
            empty_streak = 0

            for row_number, cells in enumerate(row_iterator, start=1):
                if header_row == 0:
                    if row_number > 50:
                        break
                    mapped: dict[int, str] = {}
                    observed: set[str] = set()
                    for col_number, cell in enumerate(cells, start=1):
                        normalized = normalize_header(cell.value)
                        if not normalized:
                            continue
                        canonical = lookup.get(normalized)
                        if canonical:
                            observed.add(canonical)
                            mapped[col_number] = canonical
                    if len(observed) > len(best_observed):
                        best_observed = observed
                    if not required.issubset(observed):
                        continue

                    duplicates: list[str] = []
                    for col_number, canonical in mapped.items():
                        if canonical in reverse:
                            duplicates.append(canonical)
                        else:
                            reverse[canonical] = col_number
                    if duplicates:
                        raise InputFormatError(
                            f"{label}: header duplikat ditemukan: {', '.join(sorted(set(duplicates)))}."
                        )
                    header_row = row_number
                    continue

                record: dict[str, Any] = {column: "" for column in spec["columns"]}
                nonempty = False
                for canonical, col_number in reverse.items():
                    if col_number > len(cells):
                        continue
                    cell = cells[col_number - 1]
                    value = cell.value
                    if value not in (None, ""):
                        nonempty = True
                    if canonical in spec["code_columns"]:
                        record[canonical] = normalize_code(value, getattr(cell, "number_format", None))
                    elif canonical in spec["date_columns"]:
                        record[canonical] = parse_date_value(
                            value,
                            epoch=workbook.epoch,
                            issues=issues,
                            category=label,
                            reference=f"baris {row_number}",
                            field=canonical,
                        )
                    elif isinstance(value, str):
                        record[canonical] = clean_text(value)
                    else:
                        record[canonical] = value

                if not nonempty:
                    empty_streak += 1
                    # Long blank tails are common in formatted templates.
                    if empty_streak >= 25:
                        break
                    continue
                empty_streak = 0
                key_value = record.get(spec["key"], "")
                if not keep_blank_key_rows and key_value in (None, ""):
                    issues.append(
                        Issue(
                            "WARNING",
                            label,
                            f"baris {row_number}",
                            f"Baris dilewati karena {spec['key']} kosong.",
                        )
                    )
                    continue
                record["_source_row"] = row_number
                rows.append(record)

            if header_row:
                return InputTable(
                    label=label,
                    source_name=path.name,
                    sheet_name=worksheet.title,
                    header_row=header_row,
                    rows=rows,
                )

        missing = sorted(required - best_observed)
        message = ", ".join(missing) if missing else "header tidak dikenali"
        raise InputFormatError(f"{label}: kolom wajib tidak ditemukan: {message}.")
    finally:
        workbook.close()
