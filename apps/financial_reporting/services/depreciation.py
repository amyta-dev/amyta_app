from __future__ import annotations

import calendar
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REQUIRED_FIELDS = {
    "no": ["no", "number", "nomor"],
    "description": ["description", "assetdescription", "assetname", "asset", "namaasset", "namaaset", "deskripsi"],
    "account_code": ["accountcode", "account", "kodeakun", "accountno", "coa"],
    "unit": ["unit", "qty", "quantity", "jumlahunit"],
    "purchase_date": ["purchasedate", "acquisitiondate", "tanggalperolehan", "tanggalpembelian", "datepurchased"],
    "useful_life": ["usefullifeinmonths", "usefullife", "umurmanfaatbulan", "masamanfaatbulan", "lifeinmonths"],
    "asset_amount": ["assetamount", "cost", "acquisitioncost", "amount", "nilaiasset", "nilaiaset", "harga", "hargaacquisition"],
    "depreciation_method": ["depreciationmethod", "method", "metodepenyusutan", "depmethod"],
}

OPTIONAL_FIELDS = {
    "client_name": ["clientname", "client", "namaklien", "customername"],
    "year": ["year", "period", "tahun", "reportyear"],
}

MONTH_HEADER_FORMAT = "mmm-yy"
DATE_FORMAT = "dd mmmm yyyy"
AMOUNT_FORMAT = "#,##0.00"
INTEGER_FORMAT = "#,##0"


@dataclass(frozen=True)
class AssetRecord:
    no: Any
    description: str
    account_code: Any
    unit: Any
    purchase_date: date
    useful_life: int
    asset_amount: Decimal
    depreciation_method: str


@dataclass(frozen=True)
class ReportMetadata:
    client_name: str
    year: int
    row_count: int
    source_sheet_title: str


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def make_header_map(header_values: Iterable[Any]) -> dict[str, int]:
    normalized_to_index = {normalize_header(v): idx for idx, v in enumerate(header_values, start=1) if v is not None}
    mapped: dict[str, int] = {}
    for field_name, aliases in {**REQUIRED_FIELDS, **OPTIONAL_FIELDS}.items():
        for alias in aliases:
            if alias in normalized_to_index:
                mapped[field_name] = normalized_to_index[alias]
                break
    return mapped


def find_header_row(ws) -> tuple[int, dict[str, int]]:
    max_scan_rows = min(ws.max_row, 30)
    for row_no in range(1, max_scan_rows + 1):
        values = [ws.cell(row_no, col).value for col in range(1, ws.max_column + 1)]
        header_map = make_header_map(values)
        if all(field in header_map for field in REQUIRED_FIELDS):
            return row_no, header_map

    required_labels = ", ".join([
        "No",
        "Description",
        "Account Code",
        "Unit",
        "Purchase Date",
        "Useful Life (in months)",
        "Asset Amount",
        "Depreciation Method",
    ])
    raise ValueError(f"Could not find a valid header row. Required columns: {required_labels}.")


def cell_value(ws, row_no: int, col_no: int | None) -> Any:
    if not col_no:
        return None
    return ws.cell(row_no, col_no).value


def parse_decimal(value: Any, field_name: str) -> Decimal:
    if value is None or value == "":
        raise ValueError(f"{field_name} is blank.")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))

    cleaned = str(value).strip().replace(",", "")
    cleaned = re.sub(r"[^0-9.\-]", "", cleaned)
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name} must be numeric. Value found: {value!r}") from exc


def parse_int(value: Any, field_name: str) -> int:
    number = parse_decimal(value, field_name)
    if number <= 0:
        raise ValueError(f"{field_name} must be greater than 0. Value found: {value!r}")
    return int(number)


def parse_year(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, date):
        return value.year
    match = re.search(r"(19|20)\d{2}", str(value))
    if match:
        return int(match.group(0))
    try:
        year = int(Decimal(str(value).strip()))
        if 1900 <= year <= 2100:
            return year
    except Exception:
        pass
    return None


def parse_date(value: Any, field_name: str) -> date:
    if value is None or value == "":
        raise ValueError(f"{field_name} is blank.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%m-%Y",
        "%m-%d-%Y",
        "%d %B %Y",
        "%d %b %Y",
        "%Y/%m/%d",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    try:
        return datetime.fromisoformat(text).date()
    except ValueError as exc:
        raise ValueError(f"{field_name} must be a valid date. Value found: {value!r}") from exc


def round_money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def decimal_to_float(value: Decimal) -> float:
    return float(round_money(value))


def months_before_year(purchase_date: date, year: int) -> int:
    elapsed = (year - purchase_date.year) * 12 + (1 - purchase_date.month)
    return max(0, min(elapsed, 10_000))


def month_index_from_purchase(purchase_date: date, year: int, month: int) -> int:
    return (year - purchase_date.year) * 12 + (month - purchase_date.month)


def normalize_method(method: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "", (method or "").strip().lower())
    if normalized in {"straightline", "sl", "garislurus"}:
        return "straight_line"
    if normalized in {
        "doubledeclining",
        "doubledecliningbalance",
        "doubledecliningmethod",
        "doubledecliningdepreciation",
        "ddb",
        "doublereducingbalance",
        "doubledecliningbalancemethod",
        "saldomenurunganda",
    }:
        return "double_declining"
    raise ValueError(
        f"Unsupported depreciation method: {method!r}. Supported methods are Straight Line and Double Declining Balance / DDB."
    )


def straight_line_monthly_amount(asset: AssetRecord) -> Decimal:
    return round_money(asset.asset_amount / Decimal(asset.useful_life))


def straight_line_book_value_before_year(asset: AssetRecord, year: int) -> Decimal:
    if year < asset.purchase_date.year:
        return Decimal("0.00")
    elapsed = min(months_before_year(asset.purchase_date, year), asset.useful_life)
    monthly = straight_line_monthly_amount(asset)
    depreciated = monthly * Decimal(elapsed)
    if elapsed >= asset.useful_life:
        depreciated = asset.asset_amount
    return max(Decimal("0.00"), round_money(asset.asset_amount - depreciated))


def calculate_straight_line_schedule(asset: AssetRecord, year: int) -> tuple[Decimal, list[Decimal], Decimal, Decimal]:
    if year < asset.purchase_date.year:
        zero_months = [Decimal("0.00") for _ in range(12)]
        return Decimal("0.00"), zero_months, Decimal("0.00"), Decimal("0.00")

    if year == asset.purchase_date.year:
        # This intentionally matches the provided sample output: current-year acquisitions show the asset cost as Beginning Book Value.
        beginning_book_value = round_money(asset.asset_amount)
    else:
        beginning_book_value = straight_line_book_value_before_year(asset, year)

    remaining = round_money(beginning_book_value)
    monthly_base = straight_line_monthly_amount(asset)
    monthly_amounts: list[Decimal] = []

    for month in range(1, 13):
        asset_month_index = month_index_from_purchase(asset.purchase_date, year, month)
        if 0 <= asset_month_index < asset.useful_life and remaining > 0:
            depreciation = monthly_base
            if asset_month_index == asset.useful_life - 1:
                depreciation = remaining
            depreciation = min(round_money(depreciation), remaining)
        else:
            depreciation = Decimal("0.00")
        monthly_amounts.append(depreciation)
        remaining = max(Decimal("0.00"), round_money(remaining - depreciation))

    total_depreciation = round_money(sum(monthly_amounts, Decimal("0.00")))
    ending_book_value = max(Decimal("0.00"), round_money(beginning_book_value - total_depreciation))
    return beginning_book_value, monthly_amounts, total_depreciation, ending_book_value


def active_months_for_year(asset: AssetRecord, year: int) -> list[int]:
    months: list[int] = []
    for month in range(1, 13):
        asset_month_index = month_index_from_purchase(asset.purchase_date, year, month)
        if 0 <= asset_month_index < asset.useful_life:
            months.append(month)
    return months


def annual_double_declining_depreciation(asset: AssetRecord, remaining: Decimal, year: int) -> Decimal:
    if remaining <= Decimal("0.00"):
        return Decimal("0.00")

    months = active_months_for_year(asset, year)
    if not months:
        return Decimal("0.00")

    last_asset_month_index = month_index_from_purchase(asset.purchase_date, year, months[-1])
    if last_asset_month_index >= asset.useful_life - 1:
        return round_money(remaining)

    annual_rate = Decimal("24") / Decimal(asset.useful_life)
    prorate_factor = Decimal(len(months)) / Decimal("12") if year == asset.purchase_date.year else Decimal("1.00")
    depreciation = round_money(remaining * annual_rate * prorate_factor)
    return min(depreciation, remaining)


def double_declining_book_value_before_year(asset: AssetRecord, year: int) -> Decimal:
    if year < asset.purchase_date.year:
        return Decimal("0.00")

    remaining = round_money(asset.asset_amount)
    for dep_year in range(asset.purchase_date.year, year):
        if not active_months_for_year(asset, dep_year):
            continue
        depreciation = annual_double_declining_depreciation(asset, remaining, dep_year)
        remaining = max(Decimal("0.00"), round_money(remaining - depreciation))
        if remaining <= Decimal("0.00"):
            break
    return remaining


def calculate_double_declining_schedule(asset: AssetRecord, year: int) -> tuple[Decimal, list[Decimal], Decimal, Decimal]:
    if year < asset.purchase_date.year:
        zero_months = [Decimal("0.00") for _ in range(12)]
        return Decimal("0.00"), zero_months, Decimal("0.00"), Decimal("0.00")

    if year == asset.purchase_date.year:
        beginning_book_value = round_money(asset.asset_amount)
    else:
        beginning_book_value = double_declining_book_value_before_year(asset, year)

    active_months = active_months_for_year(asset, year)
    annual_depreciation = annual_double_declining_depreciation(asset, beginning_book_value, year)

    monthly_amounts = [Decimal("0.00") for _ in range(12)]
    if active_months and annual_depreciation > Decimal("0.00"):
        unrounded_monthly = annual_depreciation / Decimal(len(active_months))
        allocated = Decimal("0.00")
        for index, month in enumerate(active_months):
            if index == len(active_months) - 1:
                monthly_dep = round_money(annual_depreciation - allocated)
            else:
                monthly_dep = round_money(unrounded_monthly)
                allocated += monthly_dep
            monthly_amounts[month - 1] = monthly_dep

    total_depreciation = round_money(sum(monthly_amounts, Decimal("0.00")))
    ending_book_value = max(Decimal("0.00"), round_money(beginning_book_value - total_depreciation))
    return beginning_book_value, monthly_amounts, total_depreciation, ending_book_value


def calculate_schedule(asset: AssetRecord, year: int) -> tuple[Decimal, list[Decimal], Decimal, Decimal]:
    method = normalize_method(asset.depreciation_method)
    if method == "double_declining":
        return calculate_double_declining_schedule(asset, year)
    return calculate_straight_line_schedule(asset, year)


def row_is_empty(ws, row_no: int, columns: Iterable[int]) -> bool:
    return all(ws.cell(row_no, col).value in (None, "") for col in columns)


def extract_assets_and_metadata(input_file: str | Path | BinaryIO) -> tuple[list[AssetRecord], str | None, int | None, str]:
    wb = load_workbook(input_file, data_only=True)
    ws = wb.active
    header_row, header_map = find_header_row(ws)

    assets: list[AssetRecord] = []
    detected_client: str | None = None
    detected_year: int | None = None
    required_cols = [header_map[field] for field in REQUIRED_FIELDS]

    for row_no in range(header_row + 1, ws.max_row + 1):
        if row_is_empty(ws, row_no, required_cols):
            continue

        if detected_client is None:
            client_value = cell_value(ws, row_no, header_map.get("client_name"))
            if client_value not in (None, ""):
                detected_client = str(client_value).strip()

        if detected_year is None:
            year_value = cell_value(ws, row_no, header_map.get("year"))
            detected_year = parse_year(year_value)

        try:
            useful_life = parse_int(cell_value(ws, row_no, header_map["useful_life"]), "Useful Life (in months)")
            asset_amount = parse_decimal(cell_value(ws, row_no, header_map["asset_amount"]), "Asset Amount")
            purchase_date = parse_date(cell_value(ws, row_no, header_map["purchase_date"]), "Purchase Date")
        except ValueError as exc:
            raise ValueError(f"Row {row_no}: {exc}") from exc

        no_value = cell_value(ws, row_no, header_map["no"])
        if isinstance(no_value, str) and no_value.startswith("="):
            no_value = len(assets) + 1

        asset = AssetRecord(
            no=no_value if no_value not in (None, "") else len(assets) + 1,
            description=str(cell_value(ws, row_no, header_map["description"]) or "").strip(),
            account_code=cell_value(ws, row_no, header_map["account_code"]),
            unit=cell_value(ws, row_no, header_map["unit"]),
            purchase_date=purchase_date,
            useful_life=useful_life,
            asset_amount=round_money(asset_amount),
            depreciation_method=str(cell_value(ws, row_no, header_map["depreciation_method"]) or "").strip(),
        )
        # Validate method early so the API can report the row with a helpful error.
        try:
            normalize_method(asset.depreciation_method)
        except ValueError as exc:
            raise ValueError(f"Row {row_no}: {exc}") from exc
        assets.append(asset)

    if not assets:
        raise ValueError("No asset rows were found below the header row.")

    return assets, detected_client, detected_year, ws.title


def apply_output_styles(ws, last_row: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="B8CCE4")
    header_font = Font(name="Arial", size=11, bold=True)
    body_font = Font(name="Arial", size=11)

    for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=23):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top")

    for cell in ws[5]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=6, max_row=last_row, min_col=1, max_col=23):
        for cell in row:
            cell.border = border
            cell.font = body_font

    for row_no in range(6, last_row + 1):
        ws.cell(row_no, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row_no, 2).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row_no, 3).alignment = Alignment(horizontal="left", vertical="top")
        ws.cell(row_no, 4).alignment = Alignment(horizontal="center", vertical="top")
        for col_no in range(5, 24):
            ws.cell(row_no, col_no).alignment = Alignment(vertical="top")

    ws.cell(2, 1).font = Font(name="Arial", size=11, bold=True)

    for row_no in range(6, last_row + 1):
        ws.cell(row_no, 5).number_format = DATE_FORMAT
        ws.cell(row_no, 6).number_format = INTEGER_FORMAT
        for col_no in [7, 9, 22, 23]:
            ws.cell(row_no, col_no).number_format = AMOUNT_FORMAT
        for col_no in range(10, 22):
            ws.cell(row_no, col_no).number_format = AMOUNT_FORMAT

    for col_no in range(10, 22):
        ws.cell(5, col_no).number_format = MONTH_HEADER_FORMAT

    widths = {
        "A": 5,
        "B": 35,
        "C": 16,
        "D": 10,
        "E": 18,
        "F": 22,
        "G": 18,
        "H": 22,
        "I": 22,
        "V": 18,
        "W": 18,
    }
    for col_no in range(10, 22):
        widths[get_column_letter(col_no)] = 13
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:W{last_row}"


def build_report_workbook(
    input_file: str | Path | BinaryIO,
    client_name: str | None = None,
    year: int | None = None,
) -> tuple[Workbook, ReportMetadata]:
    assets, detected_client, detected_year, source_sheet_title = extract_assets_and_metadata(input_file)
    report_client = (client_name or detected_client or "").strip()
    report_year = int(year or detected_year or datetime.now().year)

    output_wb = Workbook()
    ws = output_wb.active
    ws.title = source_sheet_title if source_sheet_title else "Depreciation Schedule"

    ws.cell(1, 1, "Client ")
    ws.cell(1, 3, f":  {report_client}")
    ws.cell(2, 1, "Depreciation Schedule")
    ws.cell(3, 1, "Period")
    ws.cell(3, 3, f":  {report_year}")

    headers: list[Any] = [
        "No",
        "Description",
        "Account Code",
        "Unit",
        "Purchase Date",
        "Useful Life (in months)",
        "Asset Amount",
        "Depreciation Method",
        "Beginning Book Value",
    ]
    headers.extend(date(report_year, month, 1) for month in range(1, 13))
    headers.extend(["Total Depreciation", "Ending Book Value"])

    for col_no, header in enumerate(headers, start=1):
        ws.cell(5, col_no, header)

    row_no = 6
    for asset in assets:
        beginning_bv, monthly_amounts, total_dep, ending_bv = calculate_schedule(asset, report_year)
        values: list[Any] = [
            asset.no,
            asset.description,
            str(asset.account_code) if asset.account_code is not None else "",
            asset.unit,
            asset.purchase_date,
            asset.useful_life,
            decimal_to_float(asset.asset_amount),
            asset.depreciation_method,
            decimal_to_float(beginning_bv),
        ]
        values.extend(decimal_to_float(amount) for amount in monthly_amounts)
        values.extend([decimal_to_float(total_dep), decimal_to_float(ending_bv)])
        for col_no, value in enumerate(values, start=1):
            ws.cell(row_no, col_no, value)
        row_no += 1

    last_row = row_no - 1
    apply_output_styles(ws, last_row)
    metadata = ReportMetadata(
        client_name=report_client,
        year=report_year,
        row_count=len(assets),
        source_sheet_title=source_sheet_title,
    )
    return output_wb, metadata


def build_report_bytes(
    input_file: str | Path | BinaryIO,
    client_name: str | None = None,
    year: int | None = None,
) -> tuple[BytesIO, ReportMetadata]:
    wb, metadata = build_report_workbook(input_file, client_name=client_name, year=year)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output, metadata


def save_report(
    input_path: str | Path,
    output_path: str | Path,
    client_name: str | None = None,
    year: int | None = None,
) -> ReportMetadata:
    wb, metadata = build_report_workbook(input_path, client_name=client_name, year=year)
    wb.save(output_path)
    return metadata


def safe_report_filename(client_name: str | None, year: int | None) -> str:
    safe_client = re.sub(r"[^A-Za-z0-9_.-]+", "_", (client_name or "Client")).strip("_") or "Client"
    report_year = str(year or "Report")
    return f"Depreciation Schedule - {safe_client} - {report_year}.xlsx"
