from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import ParsedStatement

_MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
_DATE_FORMAT = 'dd/mm/yyyy'
_CONTROL_CHARACTERS_RE = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')
_HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
_SUBHEADER_FILL = PatternFill("solid", fgColor="FFD9EAF7")
_WHITE_FONT = Font(color="FFFFFFFF", bold=True)
_BOLD_FONT = Font(bold=True)
_THIN_BLUE = Side(style="thin", color="FF9EBDD7")


def _excel_text(value: object, *, maximum_length: int = 32767) -> str:
    """Return XML-safe text that fits in a single Excel cell."""
    cleaned = _CONTROL_CHARACTERS_RE.sub("", str(value or ""))
    return cleaned[:maximum_length]


def build_excel(statement: ParsedStatement, source_filename: str) -> BytesIO:
    workbook = Workbook()
    transactions_sheet = workbook.active
    transactions_sheet.title = "Transactions"
    transactions_sheet.sheet_view.showGridLines = False

    headers = ["Transaction Date", "Description", "Debit", "Credit", "Balance"]
    transactions_sheet.append(headers)
    for cell in transactions_sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=_THIN_BLUE)

    for transaction in statement.transactions:
        transactions_sheet.append(
            [
                transaction.transaction_date,
                _excel_text(transaction.description),
                float(transaction.debit) if transaction.debit else None,
                float(transaction.credit) if transaction.credit else None,
                float(transaction.balance) if transaction.balance is not None else None,
            ]
        )
        row = transactions_sheet.max_row
        transactions_sheet.cell(row, 1).number_format = _DATE_FORMAT
        transactions_sheet.cell(row, 1).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        transactions_sheet.cell(row, 2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        for column in range(3, 6):
            cell = transactions_sheet.cell(row, column)
            cell.number_format = _MONEY_FORMAT
            cell.alignment = Alignment(horizontal="right", vertical="top")
        line_count = max(1, transaction.description.count("\n") + 1)
        transactions_sheet.row_dimensions[row].height = min(90, max(20, line_count * 15))

    transactions_sheet.freeze_panes = "A2"
    # Use one worksheet AutoFilter only.  A previous version added both an
    # Excel Table (which owns its own AutoFilter) and a worksheet AutoFilter
    # over the same range.  Some desktop Excel builds repair that overlapping
    # filter structure when the workbook is opened.  A standard worksheet
    # AutoFilter is sufficient for this report and produces simpler OOXML.
    transactions_sheet.auto_filter.ref = transactions_sheet.dimensions
    transactions_sheet.column_dimensions["A"].width = 19
    transactions_sheet.column_dimensions["B"].width = 72
    transactions_sheet.column_dimensions["C"].width = 20
    transactions_sheet.column_dimensions["D"].width = 20
    transactions_sheet.column_dimensions["E"].width = 22
    transactions_sheet.row_dimensions[1].height = 24
    transactions_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    transactions_sheet.page_setup.fitToWidth = 1
    transactions_sheet.page_setup.fitToHeight = 0
    transactions_sheet.print_title_rows = "1:1"

    summary = workbook.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "Bank Statement PDF Parsing - Summary"
    summary["A1"].font = Font(size=16, bold=True, color="FF1F4E78")
    summary.merge_cells("A1:B1")

    rows = [
        ("Bank", _excel_text(statement.bank_name)),
        ("Source file", _excel_text(Path(source_filename).name)),
        ("Account number", _excel_text(statement.account_number or "-")),
        ("Account name", _excel_text(statement.account_name or "-")),
        ("Currency", _excel_text(statement.currency or "-")),
        ("Period start", statement.period_start),
        ("Period end", statement.period_end),
        ("Opening balance", statement.opening_balance),
        ("Transaction count", len(statement.transactions)),
        ("Calculated total debit", statement.total_debit),
        ("Calculated total credit", statement.total_credit),
        ("Closing balance", statement.closing_balance),
        ("Reconciliation", statement.reconciliation_status()),
        ("Inferred balances", statement.inferred_balance_count),
        ("Generated at (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)),
    ]

    for row_number, (label, value) in enumerate(rows, start=3):
        summary.cell(row_number, 1, label)
        summary.cell(row_number, 2, value)
        summary.cell(row_number, 1).font = _BOLD_FONT
        summary.cell(row_number, 1).fill = _SUBHEADER_FILL
        summary.cell(row_number, 1).alignment = Alignment(vertical="top")
        summary.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")

        if label in {
            "Opening balance",
            "Calculated total debit",
            "Calculated total credit",
            "Closing balance",
        } and value is not None:
            summary.cell(row_number, 2, float(value))
            summary.cell(row_number, 2).number_format = _MONEY_FORMAT
        elif label in {"Period start", "Period end"} and value is not None:
            summary.cell(row_number, 2).number_format = _DATE_FORMAT
        elif label == "Generated at (UTC)":
            summary.cell(row_number, 2).number_format = "dd/mm/yyyy hh:mm:ss"

    warning_start = len(rows) + 5
    summary.cell(warning_start, 1, "Warnings / Notes")
    summary.cell(warning_start, 1).fill = _HEADER_FILL
    summary.cell(warning_start, 1).font = _WHITE_FONT
    summary.merge_cells(start_row=warning_start, start_column=1, end_row=warning_start, end_column=2)

    warnings = statement.warnings or ["Tidak ada warning dari parser."]
    for offset, warning in enumerate(warnings, start=1):
        summary.cell(warning_start + offset, 1, f"{offset}.")
        summary.cell(warning_start + offset, 2, _excel_text(warning))
        summary.cell(warning_start + offset, 2).alignment = Alignment(wrap_text=True)

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 68
    summary.freeze_panes = "A3"
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.print_area = f"A1:B{summary.max_row}"

    workbook.properties.title = "Parsed Bank Statement"
    workbook.properties.subject = _excel_text(statement.bank_name, maximum_length=255)
    workbook.properties.creator = "Bank Statement PDF Parser"
    workbook.properties.description = (
        "Generated from a bank statement PDF. Review warnings and reconciliation before use."
    )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def output_filename(source_filename: str, bank_code: str) -> str:
    stem = Path(source_filename).stem or "bank_statement"
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._") or "bank_statement"
    return f"{stem}_{bank_code.upper()}_parsed.xlsx"
